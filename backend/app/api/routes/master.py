from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
import httpx

from fastapi import APIRouter, Depends, File, UploadFile
from openpyxl import load_workbook
from pydantic import BaseModel, Field
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.api.deps import current_user, get_db
from app.api.errors import ApiError
from app.core.config import get_settings
from app.models.tables import Branch, Debtor, Product, ProductType, User
from app.services.audit import write_audit
from app.services.authorization import MASTER_MANAGE, has_permission, visible_branch_id

router = APIRouter(prefix="/api/v1/master", tags=["master"])


class BranchBody(BaseModel):
    code: str = Field(pattern=r"^[A-Za-z0-9_-]{1,30}$")
    name: str = Field(min_length=1, max_length=150)
    is_active: bool = True


class ProductBody(BaseModel):
    code: str = Field(pattern=r"^[A-Za-z0-9_-]{1,30}$")
    name: str = Field(min_length=1, max_length=150)
    is_active: bool = True
    business_unit: int = Field(default=0, ge=0, le=1)
    interest_rate: Decimal | None = None
    product_type_id: int | None = None


class DebtorBody(BaseModel):
    nik: str = Field(pattern=r"^[0-9]{16}$")
    full_name: str = Field(min_length=1, max_length=150)
    branch_id: int
    is_active: bool = True
    cis_id: str | None = None
    cif_id: str | None = None
    npwp: str | None = None
    birth_date: date | None = None
    birth_place: str | None = None
    mother_name: str | None = None
    gender: str | None = None
    phone: str | None = None
    address: str | None = None
    religion: str | None = None


def _clean(value: str | None) -> str | None:
    if value is None:
        return None
    value = value.strip()
    return value or None


def _debtor_profile(body: DebtorBody) -> dict:
    return {
        "cis_id": _clean(body.cis_id),
        "cif_id": _clean(body.cif_id),
        "npwp": _clean(body.npwp),
        "birth_date": body.birth_date,
        "birth_place": _clean(body.birth_place),
        "mother_name": _clean(body.mother_name),
        "gender": _clean(body.gender),
        "phone": _clean(body.phone),
        "address": _clean(body.address),
        "religion": _clean(body.religion),
    }


def _product_terms(body: ProductBody) -> dict:
    return {
        "business_unit": body.business_unit,
        "interest_rate": body.interest_rate,
        "product_type_id": body.product_type_id,
    }


def require_master_manager(user: User = Depends(current_user), db: Session = Depends(get_db)) -> User:
    if not has_permission(db, user, MASTER_MANAGE):
        raise ApiError(403, "04", "Anda tidak memiliki hak akses")
    return user


def _branch_for_write(db: Session, actor: User, branch_id: int) -> Branch:
    branch = db.get(Branch, branch_id)
    if branch is None or not branch.is_active:
        raise ApiError(400, "01", "Cabang tidak tersedia")
    scope = visible_branch_id(db, actor)
    if scope is not None and branch.id != scope:
        raise ApiError(403, "04", "Anda tidak memiliki hak akses untuk cabang ini")
    return branch


@router.get("/branches")
def list_branches(actor: User = Depends(current_user), db: Session = Depends(get_db)) -> dict:
    statement = select(Branch).order_by(Branch.name)
    scope = visible_branch_id(db, actor)
    if scope is not None:
        statement = statement.where(Branch.id == scope)
    if not has_permission(db, actor, MASTER_MANAGE):
        statement = statement.where(Branch.is_active.is_(True))
    rows = db.scalars(statement).all()
    return {
        "rcode": "00",
        "message": "Data berhasil ditampilkan",
        "result": {"items": [{"id": row.id, "code": row.code, "name": row.name, "is_active": row.is_active} for row in rows]},
    }


@router.post("/branches")
def create_branch(
    body: BranchBody,
    actor: User = Depends(require_master_manager),
    db: Session = Depends(get_db),
) -> dict:
    code = body.code.strip()
    if db.scalar(select(Branch.id).where(Branch.code == code)) is not None:
        raise ApiError(400, "01", "Kode cabang sudah digunakan")
    branch = Branch(code=code, name=body.name.strip(), is_active=body.is_active)
    db.add(branch)
    db.flush()
    write_audit(db, actor, "master.branch_created", object_type="branch", object_id=str(branch.id), after_data={"code": branch.code})
    db.commit()
    return {"rcode": "00", "message": "Cabang berhasil disimpan", "result": {"id": branch.id}}


@router.patch("/branches/{branch_id}")
def update_branch(
    branch_id: int,
    body: BranchBody,
    actor: User = Depends(require_master_manager),
    db: Session = Depends(get_db),
) -> dict:
    branch = db.get(Branch, branch_id)
    if branch is None:
        raise ApiError(404, "01", "Cabang tidak ditemukan")
    code = body.code.strip()
    taken = db.scalar(select(Branch.id).where(Branch.code == code, Branch.id != branch.id))
    if taken is not None:
        raise ApiError(400, "01", "Kode cabang sudah digunakan")
    scope = visible_branch_id(db, actor)
    if scope is not None and branch.id != scope:
        raise ApiError(403, "04", "Anda tidak memiliki hak akses untuk cabang ini")
    if branch.is_active and not body.is_active:
        active_users = db.scalar(
            select(func.count(User.id)).where(User.branch_id == branch.id, User.is_active.is_(True))
        )
        if active_users:
            raise ApiError(400, "01", "Cabang masih memiliki user aktif")
    before = {"code": branch.code, "name": branch.name, "is_active": branch.is_active}
    branch.code = code
    branch.name = body.name.strip()
    branch.is_active = body.is_active
    write_audit(
        db,
        actor,
        "master.branch_updated",
        object_type="branch",
        object_id=str(branch.id),
        before_data=before,
        after_data={"code": branch.code, "name": branch.name, "is_active": branch.is_active},
    )
    db.commit()
    return {"rcode": "00", "message": "Cabang berhasil diperbarui", "result": {"id": branch.id}}


def _plain_number(value: Decimal | None) -> str | None:
    if value is None:
        return None
    text_value = format(value, "f")
    if "." in text_value:
        text_value = text_value.rstrip("0").rstrip(".")
    return text_value


@router.get("/products")
def list_products(actor: User = Depends(current_user), db: Session = Depends(get_db)) -> dict:
    statement = (
        select(Product, ProductType.name)
        .outerjoin(ProductType, ProductType.id == Product.product_type_id)
        .order_by(Product.name)
    )
    if not has_permission(db, actor, MASTER_MANAGE):
        statement = statement.where(Product.is_active.is_(True))
    rows = db.execute(statement).all()
    types = db.scalars(select(ProductType).order_by(ProductType.name)).all()
    return {
        "rcode": "00",
        "message": "Data berhasil ditampilkan",
        "result": {
            "types": [{"id": row.id, "name": row.name} for row in types],
            "items": [
                {
                    "id": row.id,
                    "code": row.code,
                    "name": row.name,
                    "is_active": row.is_active,
                    "business_unit": row.business_unit,
                    "interest_rate": _plain_number(row.interest_rate),
                    "product_type_id": row.product_type_id,
                    "product_type_name": type_name,
                    "created_at": row.created_at.isoformat(),
                    "updated_at": row.updated_at.isoformat(),
                }
                for row, type_name in rows
            ]
        },
    }


@router.get("/products/datatables")
def list_products_datatables(
    page: int = 1,
    per_page: int = 10,
    search: str | None = None,
    code: str | None = None,
    status: str | None = None,
    branch: str | None = None,
    product_type_id: int | None = None,
    product_type: str | None = None,
    sort_by: str = "default",
    sort_dir: str = "asc",
    actor: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    is_master = has_permission(db, actor, MASTER_MANAGE)
    
    # 1. Total counts for KPI Stats
    total_all = db.scalar(select(func.count(Product.id))) or 0
    total_active = db.scalar(select(func.count(Product.id)).where(Product.is_active.is_(True))) or 0
    total_inactive = db.scalar(select(func.count(Product.id)).where(Product.is_active.is_(False))) or 0
    records_total = total_all if is_master else total_active

    # 2. Base Query with Outer Join
    query = (
        select(Product, ProductType.name)
        .outerjoin(ProductType, ProductType.id == Product.product_type_id)
    )
    if not is_master:
        query = query.where(Product.is_active.is_(True))

    # 3. Filtering
    if search and search.strip():
        kw = f"%{search.strip()}%"
        query = query.where(Product.code.ilike(kw) | Product.name.ilike(kw))

    if code and code.strip():
        query = query.where(Product.code == code.strip())

    if status and status.strip():
        st = status.strip().lower()
        if st in ("1", "aktif", "true"):
            query = query.where(Product.is_active.is_(True))
        elif st in ("0", "nonaktif", "false"):
            query = query.where(Product.is_active.is_(False))

    if branch and branch.strip():
        br = branch.strip().lower()
        if br in ("0", "konvensional", "konven"):
            query = query.where(Product.business_unit == 0)
        elif br in ("1", "syariah"):
            query = query.where(Product.business_unit == 1)

    if product_type_id:
        query = query.where(Product.product_type_id == product_type_id)
    elif product_type and product_type.strip():
        query = query.where(ProductType.name == product_type.strip())

    # Count filtered
    records_filtered = db.scalar(select(func.count()).select_from(query.subquery())) or 0

    # 4. Sorting
    sort_dir_lower = sort_dir.lower() if sort_dir else "asc"
    col_map = {
        "code": Product.code,
        "name": Product.name,
        "branch": Product.business_unit,
        "business_unit": Product.business_unit,
        "product_type": ProductType.name,
        "product_type_name": ProductType.name,
        "interest_rate": Product.interest_rate,
        "status": Product.is_active,
        "is_active": Product.is_active,
        "id": Product.id,
    }

    if sort_by in col_map:
        col = col_map[sort_by]
        query = query.order_by(col.desc().nulls_last() if sort_dir_lower == "desc" else col.asc().nulls_last())
    else:
        # Default sort: Aktif first, Konvensional first, then code
        query = query.order_by(Product.is_active.desc(), Product.business_unit.asc(), Product.code.asc())

    # 5. Pagination
    safe_page = max(1, page)
    safe_per_page = max(1, min(per_page, 500))
    offset = (safe_page - 1) * safe_per_page
    query = query.offset(offset).limit(safe_per_page)

    rows = db.execute(query).all()
    types = db.scalars(select(ProductType).order_by(ProductType.name)).all()

    # Product options for filter (Ordered: 1. is_active desc, 2. business_unit asc [konvensional first], 3. code asc)
    po_stmt = (
        select(Product.code, Product.name)
        .order_by(Product.is_active.desc(), Product.business_unit.asc(), Product.code.asc())
    )
    if not is_master:
        po_stmt = po_stmt.where(Product.is_active.is_(True))
    
    product_options = []
    seen_codes = set()
    for p_code, p_name in db.execute(po_stmt).all():
        if p_code and p_code not in seen_codes:
            seen_codes.add(p_code)
            product_options.append({"code": p_code, "name": p_name})
    
    codes = [po["code"] for po in product_options]

    items = [
        {
            "id": row.id,
            "code": row.code,
            "name": row.name,
            "is_active": row.is_active,
            "business_unit": row.business_unit,
            "interest_rate": _plain_number(row.interest_rate),
            "product_type_id": row.product_type_id,
            "product_type_name": type_name or "-",
            "created_at": row.created_at.isoformat() if row.created_at else None,
            "updated_at": row.updated_at.isoformat() if row.updated_at else None,
        }
        for row, type_name in rows
    ]

    return {
        "rcode": "00",
        "message": "Data berhasil ditampilkan",
        "result": {
            "items": items,
            "total": records_total,
            "filtered": records_filtered,
            "page": safe_page,
            "per_page": safe_per_page,
            "types": [{"id": t.id, "name": t.name} for t in types],
            "codes": list(codes),
            "product_options": product_options,
            "stats": {
                "total": total_all,
                "active": total_active,
                "inactive": total_inactive,
            },
        },
    }


@router.post("/products")
def create_product(
    body: ProductBody,
    actor: User = Depends(require_master_manager),
    db: Session = Depends(get_db),
) -> dict:
    code = body.code.strip()
    if db.scalar(select(Product.id).where(Product.code == code)) is not None:
        raise ApiError(400, "01", "Kode produk sudah digunakan")
    product = Product(code=code, name=body.name.strip(), is_active=body.is_active, **_product_terms(body))
    db.add(product)
    db.flush()
    write_audit(
        db,
        actor,
        "master.product_created",
        object_type="product",
        object_id=str(product.id),
        after_data={
            "code": product.code,
            "name": product.name,
            "business_unit": product.business_unit,
            "interest_rate": _plain_number(product.interest_rate),
            "product_type_id": product.product_type_id,
            "is_active": product.is_active,
        },
    )
    db.commit()
    return {"rcode": "00", "message": "Produk berhasil disimpan", "result": {"id": product.id}}


@router.patch("/products/{product_id}")
def update_product(
    product_id: int,
    body: ProductBody,
    actor: User = Depends(require_master_manager),
    db: Session = Depends(get_db),
) -> dict:
    product = db.get(Product, product_id)
    if product is None:
        raise ApiError(404, "01", "Produk tidak ditemukan")
    code = body.code.strip()
    taken = db.scalar(select(Product.id).where(Product.code == code, Product.id != product.id))
    if taken is not None:
        raise ApiError(400, "01", "Kode produk sudah digunakan")

    before_data = {
        "code": product.code,
        "name": product.name,
        "business_unit": product.business_unit,
        "interest_rate": _plain_number(product.interest_rate),
        "product_type_id": product.product_type_id,
        "is_active": product.is_active,
    }

    product.code = code
    product.name = body.name.strip()
    product.is_active = body.is_active
    for key, value in _product_terms(body).items():
        if key in body.model_fields_set:
            setattr(product, key, value)

    after_data = {
        "code": product.code,
        "name": product.name,
        "business_unit": product.business_unit,
        "interest_rate": _plain_number(product.interest_rate),
        "product_type_id": product.product_type_id,
        "is_active": product.is_active,
    }

    write_audit(
        db,
        actor,
        "master.product_updated",
        object_type="product",
        object_id=str(product.id),
        before_data=before_data,
        after_data=after_data,
    )
    db.commit()
    return {"rcode": "00", "message": "Produk berhasil diperbarui", "result": {"id": product.id}}


def _debtor_item(row: Debtor, branch: Branch) -> dict:
    return {
        "id": row.id,
        "nik": row.nik,
        "full_name": row.full_name,
        "branch_id": row.branch_id,
        "branch_code": branch.code,
        "branch_name": branch.name,
        "is_active": row.is_active,
        "cis_id": row.cis_id,
        "cif_id": row.cif_id,
        "npwp": row.npwp,
        "birth_date": row.birth_date.isoformat() if row.birth_date else None,
        "birth_place": row.birth_place,
        "mother_name": row.mother_name,
        "gender": row.gender,
        "phone": row.phone,
        "address": row.address,
        "religion": row.religion,
    }


@router.get("/debtors")
def list_debtors(actor: User = Depends(current_user), db: Session = Depends(get_db)) -> dict:
    statement = select(Debtor, Branch).join(Branch, Branch.id == Debtor.branch_id).order_by(Debtor.full_name)
    scope = visible_branch_id(db, actor)
    if scope is not None:
        statement = statement.where(Debtor.branch_id == scope)
    if not has_permission(db, actor, MASTER_MANAGE):
        statement = statement.where(Debtor.is_active.is_(True))
    items = [_debtor_item(debtor, branch) for debtor, branch in db.execute(statement).all()]
    return {"rcode": "00", "message": "Data berhasil ditampilkan", "result": {"items": items}}


def generate_cis_id(db: Session) -> str:
    """
    Format CIS ID: 10 digit
    - 2 digit pertama: 02 (perorangan)
    - 2 digit setelahnya: tahun saat register (contoh: 2026 -> 26)
    - 6 digit sisanya: nomor urut (zero-padded)
    """
    prefix = f"02{date.today().strftime('%y')}"
    max_cis = db.scalar(
        select(func.max(Debtor.cis_id)).where(
            Debtor.cis_id.like(f"{prefix}%"),
            func.length(Debtor.cis_id) == 10,
        )
    )
    if max_cis and len(max_cis) == 10 and max_cis[4:].isdigit():
        next_seq = int(max_cis[4:]) + 1
    else:
        next_seq = 1
    return f"{prefix}{next_seq:06d}"


@router.get("/debtors/next-cis-id")
def get_next_cis_id(actor: User = Depends(current_user), db: Session = Depends(get_db)) -> dict:
    cis_id = generate_cis_id(db)
    return {"rcode": "00", "message": "OK", "result": {"cis_id": cis_id}}


@router.get("/debtors/inquiry-cif/{cif_id}")
def inquiry_cif(
    cif_id: str,
    actor: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    clean_cif = cif_id.strip()
    if not clean_cif:
        raise ApiError(400, "01", "Nomor CIF wajib diisi")

    # 1. Cek database lokal terlebih dahulu
    statement = select(Debtor, Branch).join(Branch, Branch.id == Debtor.branch_id).where(Debtor.cif_id == clean_cif)
    result = db.execute(statement).first()
    if result is not None:
        debtor, branch = result
        return {
            "rcode": "00",
            "message": "Data nasabah ditemukan di database lokal",
            "result": {
                "found": True,
                "source": "local",
                "debtor": _debtor_item(debtor, branch),
            },
        }

    # 2. Fallback ke API Service Gateway (inqCIFByCIF — reqid: 00035)
    settings = get_settings()
    gateway_data = _call_gateway_cif_inquiry(clean_cif, settings)
    if gateway_data is not None:
        return {
            "rcode": "00",
            "message": "Data nasabah ditemukan di Core Banking (Gateway)",
            "result": {
                "found": True,
                "source": "gateway",
                "debtor": gateway_data,
            },
        }

    return {
        "rcode": "00",
        "message": f"CIF {clean_cif} belum terdaftar di sistem lokal maupun Core Banking.",
        "result": {
            "found": False,
            "cif_id": clean_cif,
        },
    }


def _call_gateway_cif_inquiry(cif_id: str, settings) -> dict | None:
    """
    Panggil Gateway /inqCIFByCIF (reqid: 00035) untuk mengambil data nasabah
    dari Core Banking berdasarkan CIF ID.
    Kembalikan dict profile yang compatible dengan _debtor_item atau None jika gagal/tidak ditemukan.
    """
    host = settings.gateway_host_url.rstrip("/")
    if not host or host == "http://change-me-gateway-host":
        # Gateway belum dikonfigurasi — skip tanpa error
        return None

    url = f"{host}/gateway/inqCIFByCIF"
    payload = {
        "reqid": "00035",
        "channelId": settings.gateway_channel_id,
        "userGtw": settings.gateway_user_gtw,
        "cifid": cif_id,
        "-": "",
    }
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    headers = {
        "Content-Type": "application/json",
        "X-Api-Key": settings.gateway_api_key,
        "X-Client-Id": settings.gateway_client_id,
        "X-Client-Secret": settings.gateway_client_secret,
        "X-Timestamp": timestamp,
        "X-Signature": settings.gateway_signature or "CABUinapmSdD1j8lIqo8Qvvc+ksRK2g2qDEUlwUfzCM=",
    }

    try:
        resp = httpx.post(url, json=payload, headers=headers, timeout=10)
        data = resp.json()
    except Exception:
        return None

    if not isinstance(data, dict):
        return None

    # Periksa response code gateway (biasanya rcode "00" = sukses)
    rcode = data.get("rcode") or data.get("statusId") or data.get("responseCode")
    result = data.get("result") or data.get("data") or {}

    cif_info = result if isinstance(result, dict) else {}

    def _get(*keys):
        for k in keys:
            v = cif_info.get(k) or cif_info.get(k.lower()) or cif_info.get(k.upper())
            if v is not None and str(v).strip():
                return str(v).strip()
        return None

    def _format_date(d_str: str | None) -> str | None:
        if not d_str:
            return None
        clean = d_str.strip()
        # YYYYMMDD -> YYYY-MM-DD
        if len(clean) == 8 and clean.isdigit():
            return f"{clean[:4]}-{clean[4:6]}-{clean[6:]}"
        # DD-MM-YYYY or DD/MM/YYYY -> YYYY-MM-DD
        if len(clean) == 10 and (clean[2] in "-/" and clean[5] in "-/"):
            return f"{clean[6:]}-{clean[3:5]}-{clean[:2]}"
        # If ISO like YYYY-MM-DD
        if len(clean) >= 10 and clean[4] == "-" and clean[7] == "-":
            return clean[:10]
        return clean

    def _format_gender(g_str: str | None) -> str | None:
        if not g_str:
            return None
        upper = g_str.strip().upper()
        if upper in ("L", "M", "LAKI-LAKI", "PRIA", "MALE", "1"):
            return "Laki-laki"
        if upper in ("P", "F", "PEREMPUAN", "WANITA", "FEMALE", "2"):
            return "Perempuan"
        return g_str.strip().title()

    def _format_religion(r_str: str | None) -> str | None:
        if not r_str:
            return None
        upper = r_str.strip().upper()
        rel_map = {
            "1": "Islam", "ISLAM": "Islam",
            "2": "Protestan", "KRISTEN": "Protestan", "PROTESTAN": "Protestan", "KRISTEN PROTESTAN": "Protestan",
            "3": "Katolik", "KATOLIK": "Katolik", "KRISTEN KATOLIK": "Katolik",
            "4": "Hindu", "HINDU": "Hindu",
            "5": "Buddha", "BUDHA": "Buddha", "BUDDHA": "Buddha",
            "6": "Konghucu", "KONGHUCU": "Konghucu", "KHONGHUCU": "Konghucu",
            "7": "Kepercayaan", "KEPERCAYAAN": "Kepercayaan",
        }
        return rel_map.get(upper, r_str.strip().title())

    full_name = _get("CIFNM", "nama", "name", "fullName", "cifnm", "CUST_NAME", "NAMA_NASABAH")
    nik = _get("IDNBR", "idnbr", "nik", "ktpNo", "NO_IDENTITAS", "KTP")
    birth_date = _format_date(_get("BRTDT", "brtdt", "birthDate", "tglLahir", "TGL_LAHIR"))
    birth_place = _get("BRTPLC", "brtplc", "birthPlace", "tmptLahir", "TEMPAT_LAHIR")
    phone = _get("HP", "PHONE", "hp", "phone", "telepon", "mobile", "NO_HP", "TELP")
    address = _get("ADDR", "addr", "address", "alamat", "ADDRESS", "ALAMAT")
    mother_name = _get("MOTHRNM", "mothrnm", "motherName", "ibuKandung", "NAMA_IBU")
    npwp = _get("NPWP", "npwp", "NO_NPWP")
    gender = _format_gender(_get("GENDER", "gender", "jenisKelamin", "JENIS_KELAMIN"))
    religion = _format_religion(_get("RELIGION", "religion", "agama", "AGAMA"))

    if not full_name and str(rcode) not in ("00", "0", "1"):
        return None

    return {
        "id": None,
        "nik": nik,
        "full_name": full_name,
        "cif_id": cif_id,
        "cis_id": None,
        "branch_id": None,
        "branch_code": None,
        "branch_name": None,
        "is_active": True,
        "birth_date": birth_date,
        "birth_place": birth_place,
        "phone": phone,
        "gender": gender,
        "religion": religion,
        "npwp": npwp,
        "address": address,
        "mother_name": mother_name,
        "_gateway_raw": cif_info,
    }



@router.post("/debtors")
def create_debtor(
    body: DebtorBody,
    actor: User = Depends(require_master_manager),
    db: Session = Depends(get_db),
) -> dict:
    if db.scalar(select(Debtor.id).where(Debtor.nik == body.nik)) is not None:
        raise ApiError(400, "01", "NIK sudah terdaftar")
    _branch_for_write(db, actor, body.branch_id)
    profile = _debtor_profile(body)
    # CIS ID create by sistem (10 digit: 02 + YY + 6-digit sequence)
    if not profile.get("cis_id") or len(str(profile.get("cis_id"))) != 10:
        profile["cis_id"] = generate_cis_id(db)

    debtor = Debtor(
        nik=body.nik,
        full_name=body.full_name.strip(),
        branch_id=body.branch_id,
        is_active=body.is_active,
        created_by=actor.id,
        **profile,
    )
    db.add(debtor)
    db.flush()
    write_audit(db, actor, "master.debtor_created", object_type="debtor", object_id=str(debtor.id), after_data={"nik": debtor.nik, "cis_id": debtor.cis_id})
    db.commit()
    return {"rcode": "00", "message": "Debitur berhasil disimpan", "result": {"id": debtor.id, "cis_id": debtor.cis_id}}


@router.patch("/debtors/{debtor_id}")
def update_debtor(
    debtor_id: int,
    body: DebtorBody,
    actor: User = Depends(require_master_manager),
    db: Session = Depends(get_db),
) -> dict:
    debtor = db.get(Debtor, debtor_id)
    if debtor is None:
        raise ApiError(404, "01", "Debitur tidak ditemukan")
    scope = visible_branch_id(db, actor)
    if scope is not None and debtor.branch_id != scope:
        raise ApiError(403, "04", "Anda tidak memiliki hak akses untuk cabang ini")
    taken = db.scalar(select(Debtor.id).where(Debtor.nik == body.nik, Debtor.id != debtor.id))
    if taken is not None:
        raise ApiError(400, "01", "NIK sudah terdaftar")
    _branch_for_write(db, actor, body.branch_id)
    before = {
        "nik": debtor.nik,
        "full_name": debtor.full_name,
        "branch_id": debtor.branch_id,
        "is_active": debtor.is_active,
    }
    debtor.nik = body.nik
    debtor.full_name = body.full_name.strip()
    debtor.branch_id = body.branch_id
    debtor.is_active = body.is_active
    for key, value in _debtor_profile(body).items():
        if key in body.model_fields_set:
            setattr(debtor, key, value)
    write_audit(
        db,
        actor,
        "master.debtor_updated",
        object_type="debtor",
        object_id=str(debtor.id),
        before_data=before,
        after_data={"nik": debtor.nik, "full_name": debtor.full_name, "branch_id": debtor.branch_id, "is_active": debtor.is_active},
    )
    db.commit()
    return {"rcode": "00", "message": "Debitur berhasil diperbarui", "result": {"id": debtor.id}}


def _cell(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


@router.post("/debtors/import")
async def import_debtors(
    file: UploadFile = File(...),
    actor: User = Depends(require_master_manager),
    db: Session = Depends(get_db),
) -> dict:
    filename = (file.filename or "").lower()
    if not filename.endswith(".xlsx"):
        raise ApiError(400, "01", "File harus berformat .xlsx")
    content = await file.read()
    if len(content) > 5 * 1024 * 1024:
        raise ApiError(400, "01", "Ukuran file melebihi 5 MB")
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
    except Exception as exc:
        raise ApiError(400, "01", "File Excel tidak dapat dibaca") from exc
    if not rows:
        raise ApiError(400, "01", "File Excel kosong")
    if len(rows) > 2001:
        raise ApiError(400, "01", "Jumlah baris melebihi 2000")
    headers = [_cell(item).lower() for item in rows[0]]
    if "nik" not in headers or "nama" not in headers:
        raise ApiError(400, "01", "Kolom wajib: NIK dan Nama")
    nik_index = headers.index("nik")
    name_index = headers.index("nama")
    created = 0
    updated = 0
    rejected: list[dict] = []
    for offset, row in enumerate(rows[1:], start=2):
        nik = _cell(row[nik_index] if nik_index < len(row) else "")
        name = _cell(row[name_index] if name_index < len(row) else "")
        if nik == "" and name == "":
            continue
        if len(nik) != 16 or not nik.isdigit() or name == "":
            rejected.append({"row": offset, "nik": nik, "reason": "NIK harus 16 digit dan nama wajib diisi"})
            continue
        existing = db.scalar(select(Debtor).where(Debtor.nik == nik))
        scope = visible_branch_id(db, actor)
        if existing is None:
            target_branch = scope if scope is not None else actor.branch_id
            debtor = Debtor(nik=nik, full_name=name, branch_id=target_branch, is_active=True, created_by=actor.id)
            db.add(debtor)
            created += 1
            continue
        if scope is not None and existing.branch_id != scope:
            rejected.append({"row": offset, "nik": nik, "reason": "NIK berada di luar cabang Anda"})
            continue
        existing.full_name = name
        updated += 1
    write_audit(
        db,
        actor,
        "master.debtor_imported",
        object_type="debtor",
        object_id="import",
        after_data={"created": created, "updated": updated, "rejected": len(rejected)},
    )
    db.commit()
    return {
        "rcode": "00",
        "message": "Import debitur selesai",
        "result": {"created": created, "updated": updated, "rejected": rejected},
    }
